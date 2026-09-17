#!/usr/bin/env python3
"""Assemble the supplied theory, introduction and experiments without replacing source files."""
from __future__ import annotations
import argparse, copy, hashlib, json, re, subprocess, tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

TITLE = '语义引导高斯时序定位的视听广义零样本学习'
ABSTRACT = ('视听广义零样本学习通过建立音视频表示与类别语义之间的关联，实现对已见类别与未见类别的联合识别。'
'现有方法通常将类别语义限制在最终匹配阶段，难以从长视频中聚焦局部判别证据，同时受到类别原型可分性不足与视听模态干扰的影响。'
'针对上述问题，本文提出语义引导高斯时序定位框架 SGTG-AVGZSL，将类别语义作为条件信号贯穿时序编码与模态融合。'
'首先，判别性语义原型优化模块在增强类别分离度的同时保持语义相对排序，为已见与未见类别建立可迁移的条件原型。'
'其次，语义引导高斯时序定位模块利用类别原型分别预测视觉与音频的多高斯时间分布，并通过中心锚定和自适应路由聚合类别相关片段。'
'随后，解耦模态交互与门控融合模块分别建模模态内与模态间关系，以类别条件三路门控融合两类关系表示及直接时序证据。'
'在 VGGSound-GZSL、UCF-GZSL 和 ActivityNet-GZSL 三个基准上的对比结果中，本文方法的调和均值分别为 17.62%、61.28% 和 31.44%，'
'相较 ClipClap-GZSL 分别提高 1.89、6.15 和 4.33 个百分点。')
KEYWORDS = '视听广义零样本学习；类别语义原型；高斯时序定位；解耦模态交互；动态门控融合'
CONCLUSION = ('本文提出了语义引导高斯时序定位的视听广义零样本学习框架 SGTG-AVGZSL，将类别语义由最终匹配阶段的静态参照扩展为贯穿编码过程的条件信号。'
'框架首先通过判别性语义原型优化兼顾类别可分性与语义结构保持，再以优化后的原型驱动视觉和音频的独立多高斯时序定位，'
'最后利用解耦模态交互与类别条件三路门控组织和融合筛选后的证据。三个模块分别作用于语义查询、时序证据定位与多模态证据组织，'
'共同构成面向已见与未见类别的统一识别过程。三个基准上的对比结果体现了该框架在已见类别识别与未见类别泛化之间的综合平衡。'
'后续研究将进一步探索候选类别条件化计算的共享与压缩，以及复杂噪声和更开放类别空间下的时序证据建模。')


def xp(e, path):
    return etree.ElementTree(e).xpath(path, namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','v':'urn:schemas-microsoft-com:vml'})


def element(tag: str, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items(): e.set(qn('w:'+k), str(v))
    return e


def normal_words(s: str) -> str:
    return s.replace('零次学习', '零样本学习').replace('音视频广义零样本学习', '视听广义零样本学习')


def replace_spanning(p, old, new):
    nodes=p._p.xpath('.//w:t')
    joined=''.join(t.text or '' for t in nodes)
    offsets=[];pos=0
    for t in nodes: offsets.append(pos);pos+=len(t.text or '')
    for match in reversed(list(re.finditer(re.escape(old),joined))):
        start,end=match.span()
        used=[i for i,t in enumerate(nodes) if offsets[i]<end and offsets[i]+len(t.text or '')>start]
        if not used:continue
        first,last=used[0],used[-1]
        head=(nodes[first].text or '')[:start-offsets[first]]
        tail=(nodes[last].text or '')[end-offsets[last]:]
        nodes[first].text=head+new+(tail if first==last else '')
        for i in used[1:-1]:nodes[i].text=''
        if last!=first:nodes[last].text=tail


def run_font(run, size=10.5, east='宋体', latin='Times New Roman'):
    run.font.name = latin; run.font.size = Pt(size)
    run.font.color.rgb = __import__('docx.shared', fromlist=['RGBColor']).RGBColor(0,0,0)
    rp = run._r.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None: rf = element('w:rFonts'); rp.insert(0,rf)
    for k,v in [('ascii',latin),('hAnsi',latin),('cs',latin),('eastAsia',east)]: rf.set(qn('w:'+k),v)
    for child in list(rp):
        if child.tag in (qn('w:highlight'),qn('w:shd')): rp.remove(child)


def format_para(p, kind='body', size=None):
    pf=p.paragraph_format
    pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5
    pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(21)
    pf.keep_with_next=False;pf.keep_together=False;pf.widow_control=True;pf.page_break_before=False
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    pr=p._p.get_or_add_pPr()
    for c in list(pr):
        if c.tag in [qn('w:numPr'),qn('w:tabs'),qn('w:shd'),qn('w:sectPr')]:pr.remove(c)
    ind=pr.find(qn('w:ind'))
    if ind is not None:
        for a in list(ind.attrib):
            if a.rsplit('}',1)[-1].endswith('Chars'): del ind.attrib[a]
    replace_spanning(p,'零次学习','零样本学习')
    replace_spanning(p,'音视频广义零样本学习','视听广义零样本学习')
    fs=size or 10.5
    if kind.startswith('h') and kind[1:].isdigit():
        lev=int(kind[1:]); p.style=f'Heading {lev}';fs={1:18,2:16,3:12}[lev]
        pf.first_line_indent=Pt(0);pf.space_before=Pt(10 if lev<3 else 7);pf.space_after=Pt(5)
        pf.line_spacing=1.15;pf.keep_with_next=True;pf.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.style='Normal'
    if kind=='title':
        p.style='Title';fs=18;pf.first_line_indent=Pt(0);pf.space_after=Pt(10);pf.line_spacing=1.2;pf.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if kind in ('caption','figure_caption','note'):
        fs=10 if kind!='note' else 9
        pf.first_line_indent=Pt(0);pf.line_spacing=1.2;pf.space_before=Pt(4);pf.space_after=Pt(4);pf.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if kind=='caption' else WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.keep_with_next=kind=='caption'
    if kind in ('abstract','keywords'):
        pf.first_line_indent=Pt(0);pf.line_spacing=1.3;pf.space_after=Pt(5)
    if kind=='reference':
        fs=9.5;pf.first_line_indent=Pt(-22);pf.left_indent=Pt(22);pf.line_spacing=1.15;pf.space_after=Pt(4);pf.keep_together=True
    if kind=='equation':
        fs=size or 10.5;pf.first_line_indent=Pt(0);pf.line_spacing=1.0;pf.space_before=Pt(5);pf.space_after=Pt(5);pf.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        run_font(r,fs,east='黑体' if kind.startswith('h') or kind=='title' else '宋体')
        if kind.startswith('h') or kind=='title':r.bold=True
    for rp in p._p.xpath('.//m:r/w:rPr'):
        for tag,value in [('w:sz',str(int(fs*2))),('w:szCs',str(int(fs*2))),('w:color','000000')]:
            ee=rp.find(qn(tag))
            if ee is None:ee=OxmlElement(tag);rp.append(ee)
            ee.set(qn('w:val'),value)
    return p


def setup_document(d, reference):
    sec=d.sections[0]; rs=reference.sections[0]
    sec.page_width=rs.page_width;sec.page_height=rs.page_height
    sec.top_margin=rs.top_margin;sec.bottom_margin=rs.bottom_margin
    sec.left_margin=rs.left_margin;sec.right_margin=rs.right_margin
    sec.header_distance=Cm(1.2);sec.footer_distance=Cm(1.2)
    for sn in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        if sn not in d.styles:d.styles.add_style(sn,WD_STYLE_TYPE.PARAGRAPH)
        s=d.styles[sn];s.font.name='Times New Roman';s.font.size=Pt(10.5)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        if s.element.pPr is not None:
            for c in list(s.element.pPr):
                if c.tag in [qn('w:numPr'),qn('w:pageBreakBefore')]:s.element.pPr.remove(c)
    for header in [sec.header,sec.first_page_header,sec.even_page_header]:
        for p in header.paragraphs:p.clear()
    for footer in [sec.footer,sec.first_page_footer,sec.even_page_footer]:
        for p in footer.paragraphs:p.clear()
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    r=p.add_run();run_font(r,9)
    fld=element('w:fldSimple',instr='PAGE');fr=element('w:r');ft=element('w:t');ft.text='1';fr.append(ft);fld.append(fr);p._p.append(fld)
    d.core_properties.title=TITLE;d.core_properties.subject='SGTG-AVGZSL';d.core_properties.comments=''
    return (sec.page_width-sec.left_margin-sec.right_margin)/12700


def clean_numbered_equation(p, width):
    om=p._p.xpath('./m:oMathPara/m:oMath')
    if not om:return None
    original=om[0];arr=original.find(qn('m:eqArr'))
    if arr is None:return None
    es=arr.findall(qn('m:e'))
    if len(es)!=1:return None
    content=es[0];children=list(content)
    sep=next((i for i,c in enumerate(children) if ''.join(xp(c,'.//m:t/text()'))=='#'),None)
    if sep is None:return None
    n=int(''.join(xp(children[sep+1],'.//m:t/text()')))
    newom=OxmlElement('m:oMath')
    for c in children[:sep]: newom.append(copy.deepcopy(c))
    for t in xp(newom,'.//m:t'):
        if t.text and '&' in t.text:t.text=t.text.replace('&','')
    p.clear();format_para(p,'equation',9.5 if n in [27] else 10.5)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(width/2),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(width),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t');p._p.append(newom);r=p.add_run(f'\t（{n}）');run_font(r,10.5)
    for rp in xp(newom,'.//m:r/w:rPr'):
        for tag in ['w:sz','w:szCs']:
            ee=rp.find(qn(tag))
            if ee is None:ee=OxmlElement(tag);rp.append(ee)
            ee.set(qn('w:val'),'19' if n==27 else '21')
    return n


def source_refs(s):
    return {int(n):v.strip() for n,v in re.findall(r'^\[(\d+)\]\s*(.*)$',s,re.M)}


def prepared_markdown(src):
    intro=(src/'引言与相关工作.md').read_text(encoding='utf-8')
    exp=(src/'实验部分.md').read_text(encoding='utf-8')
    refs=source_refs(intro);refs.update(source_refs(exp))
    intro=re.split(r'^# 参考文献\s*$',intro,flags=re.M)[0].strip()
    exp=re.split(r'^\*\*参考文献\*\*\s*$',exp,flags=re.M)[0].strip().rstrip('-').strip()
    intro=intro.replace('的策略，如图 1 所示：','的策略：').replace('（图 1(a)）','').replace('（图 1(b)）','')
    intro=intro.replace('据我们所知，这是首个面向 AVGZSL 的语义条件连续时序定位机制。','由此，类别语义不仅参与最终匹配，还直接作用于时序证据的定位与聚合。')
    for ds in ['VGGSound','UCF','ActivityNet']:
        intro=intro.replace(ds+'-GZSLcls',ds+'-GZSL$^{\\mathrm{cls}}$')
    exp=exp.replace('每段包含 32 帧，音频取对应的时间区间。','每个片段包含一帧代表图像与对应时段的音频波形。')
    exp=exp.replace('以及式（6）的联合目标 $\\mathcal{L}_{proto}=(1-\\alpha)\\mathcal{L}_{rank}+\\alpha\\mathcal{L}_{sep}$。','以及式（6）定义的类别分离与语义保持联合优化目标。')
    exp=exp.replace('联合目标<br>ℒ<sub>proto</sub>（式（6））','联合目标<br>ℒ（式（6））')
    prefix,tail=exp.split('## 4.2 与先进方法的对比',1)
    data=prefix.split('**数据集。**',1)[1].split('**评价指标。**',1)[0].strip()
    metrics=prefix.split('**评价指标。**',1)[1].split('**实现细节。**',1)[0].strip()
    implementation=prefix.split('**实现细节。**',1)[1].split('**训练与评估协议。**',1)[0].strip()
    protocol=prefix.split('**训练与评估协议。**',1)[1].split('**对比方法。**',1)[0].strip()
    comparison=prefix.split('**对比方法。**',1)[1].strip()
    exp=('# 4 实验\n\n## 4.1 实验设置\n\n### 4.1.1 数据集\n\n'+data+
         '\n\n### 4.1.2 实现细节\n\n'+implementation+'\n\n**训练与评估协议。**'+protocol+
         '\n\n### 4.1.3 评价指标\n\n'+metrics+'\n\n### 4.1.4 对比方法\n\n'+comparison+
         '\n\n## 4.2 定量分析'+tail)
    exp=exp.replace('<div align="right">（33）</div>','WH_EQUATION_NUMBER_33')
    table_html=[]
    def table_slot(m):
        table_html.append(m.group(0));return '\n\nWH_TABLE_'+str(len(table_html)).zfill(2)+'\n\n'
    exp=re.sub(r'<table>.*?</table>',table_slot,exp,flags=re.S)
    assert len(table_html)==9
    order=[]
    for text in [intro,'[9]',exp,'\n'.join(table_html)]:
        for token in re.findall(r'\[(\d+(?:,\s*\d+)*)\]',text):
            for n in [int(x) for x in token.split(',')]:
                if n in refs and n not in order:order.append(n)
    for n in sorted(refs):
        if n not in order:order.append(n)
    mapping={old:i+1 for i,old in enumerate(order)}
    def renumber(s):
        return re.sub(r'\[(\d+(?:,\s*\d+)*)\]',lambda m:'['+','.join(str(mapping.get(int(x),int(x))) for x in m[1].split(','))+']',s)
    return normal_words(renumber(intro)),normal_words(renumber(exp)),[renumber(x) for x in table_html],refs,order,mapping


def html_content(p,node,attrs=None):
    attrs=attrs or {}
    if isinstance(node,NavigableString):
        if not str(node):return
        r=p.add_run(str(node));run_font(r,8)
        for k,v in attrs.items():setattr(r.font,k,v)
        return
    if not isinstance(node,Tag):return
    a=dict(attrs)
    if node.name in ('strong','b'):a['bold']=True
    if node.name in ('i','em'):a['italic']=True
    if node.name in ('ins','u'):a['underline']=True
    if node.name=='sup':a['superscript']=True
    if node.name=='sub':a['subscript']=True
    if node.name=='br':p.add_run().add_break();return
    for child in node.children:html_content(p,child,a)


def set_cell_border(cell,where,val='single',sz='6'):
    pr=cell._tc.get_or_add_tcPr();bs=pr.find(qn('w:tcBorders'))
    if bs is None:bs=element('w:tcBorders');pr.append(bs)
    e=bs.find(qn('w:'+where))
    if e is None:e=element('w:'+where);bs.append(e)
    for k,v in [('val',val),('sz',sz),('color','000000')]:e.set(qn('w:'+k),v)


def native_table(doc,html,index,width):
    soup=BeautifulSoup(html,'html.parser');rows=soup.find_all('tr')
    occupied={};items=[];maxcol=0
    for ri,row in enumerate(rows):
        ci=0
        for cell in row.find_all(['th','td'],recursive=False):
            while (ri,ci) in occupied:ci+=1
            rs=int(cell.get('rowspan',1));cs=int(cell.get('colspan',1))
            items.append((ri,ci,rs,cs,cell))
            for rr in range(ri,ri+rs):
                for cc in range(ci,ci+cs):occupied[(rr,cc)]=True
            ci+=cs;maxcol=max(maxcol,ci)
    t=doc.add_table(rows=len(rows),cols=maxcol);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tp=t._tbl.tblPr
    for x in list(tp):
        if x.tag in [qn('w:tblBorders'),qn('w:tblInd'),qn('w:tblCellMar')]:tp.remove(x)
    tw=tp.find(qn('w:tblW'));tw.set(qn('w:w'),str(round(width*20)));tw.set(qn('w:type'),'dxa')
    borders=element('w:tblBorders')
    for b in ['top','left','bottom','right','insideH','insideV']:borders.append(element('w:'+b,val='nil'))
    tp.append(borders)
    mar=element('w:tblCellMar')
    for k,v in [('top',35),('bottom',35),('left',15),('right',15)]:mar.append(element('w:'+k,w=v,type='dxa'))
    tp.append(mar)
    if index==1:widths=[112,30,31,31]+[(width-204)/6]*6
    elif index==3:widths=[27]*3+[(width-81)/12]*12
    else:
        first={2:77,4:101,5:29,6:100,7:101,8:78,9:91}[index]
        widths=[first]+[(width-first)/12]*12
    assert len(widths)==maxcol
    for col,ww in zip(t.columns,widths):col.width=Pt(ww)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(element('w:cantSplit'))
        for ci,c in enumerate(row.cells):c.width=Pt(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ri,ci,rs,cs,node in items:
        cell=t.cell(ri,ci)
        if rs>1 or cs>1:cell=cell.merge(t.cell(ri+rs-1,ci+cs-1))
        cell.text='';p=cell.paragraphs[0]
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(0);p.paragraph_format.right_indent=Pt(0)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
        p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=(ri<len(rows)-1)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if node.get('align')=='left' else WD_ALIGN_PARAGRAPH.CENTER
        for child in node.children:html_content(p,child)
        if node.name=='th':
            for r in p.runs:r.bold=True
        if ri==0:set_cell_border(cell,'top',sz='10')
        if ri+rs==len(rows):set_cell_border(cell,'bottom',sz='10')
    nhead=len(soup.find('thead').find_all('tr'))
    for row in t.rows[:nhead]:row._tr.get_or_add_trPr().append(element('w:tblHeader'))
    for ci in range(maxcol):set_cell_border(t.cell(nhead-1,ci),'bottom',sz='6')
    for ri,ci,rs,cs,node in items:
        if ri==0 and cs>1:set_cell_border(t.cell(ri,ci),'bottom',sz='4')
    return t


def convert_md(md,work,name):
    md=md.replace('$^{\\mathrm{cls}}$','^cls^')
    source=work/(name+'.md');out=work/(name+'.docx');source.write_text(md,encoding='utf-8')
    subprocess.run(['pandoc',str(source),'-f','markdown+tex_math_dollars+raw_html','-t','docx','--standalone','-o',str(out)],check=True)
    return Document(out)


def prepare_part(part,d,table_html,width,is_exp=False):
    nodes=[];last=None;contribution=0
    for pp in part.paragraphs:
        text=pp.text.strip()
        if text.startswith('WH_TABLE_'):
            idx=int(text.rsplit('_',1)[1]);table=native_table(d,table_html[idx-1],idx,width)
            table._tbl.getparent().remove(table._tbl);nodes.append(table._tbl);last=None;continue
        if text=='WH_EQUATION_NUMBER_33':
            if last is None:raise ValueError('Metric equation marker lost')
            om=last._p.xpath('.//m:oMath')[0];om=copy.deepcopy(om);last.clear();format_para(last,'equation')
            last.paragraph_format.tab_stops.add_tab_stop(Pt(width/2),WD_TAB_ALIGNMENT.CENTER)
            last.paragraph_format.tab_stops.add_tab_stop(Pt(width),WD_TAB_ALIGNMENT.RIGHT)
            last.add_run('\t');last._p.append(om);run_font(last.add_run('\t（33）'),10.5);continue
        p=Paragraph(copy.deepcopy(pp._p),d._body)
        style=pp.style.name
        if style.startswith('Heading'):
            kind='h'+style.split()[-1]
        elif (re.match(r'^表\s*\d+',text) and text.endswith('（%）')) or text=='表 1 三个基准数据集的类别划分':kind='caption'
        elif text.startswith('注：'):kind='note'
        else:kind='body'
        format_para(p,kind)
        if style in ('Compact','List Paragraph') or pp._p.xpath('./w:pPr/w:numPr'):
            p.paragraph_format.left_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(21)
            if not is_exp:
                contribution+=1
                rr=p.add_run(f'（{contribution}）');run_font(rr);p._p.insert(1,rr._r)
        nodes.append(p._p);last=p
    return nodes


def stable_zip(path):
    with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    with ZipFile(path,'w',ZIP_DEFLATED,compresslevel=9) as z:
        for n in sorted(parts):
            zi=ZipInfo(n,(2026,9,17,0,0,0));zi.compress_type=ZIP_DEFLATED;zi.external_attr=0o600<<16
            z.writestr(zi,parts[n])


def build(src,out,work):
    src=Path(src);out=Path(out);work=Path(work);out.parent.mkdir(parents=True,exist_ok=True);work.mkdir(parents=True,exist_ok=True)
    d=Document(src/'王虎-理论部分-0908.docx');ref=Document(src/'KA-GZAL中文.docx')
    original_paras=list(d.paragraphs);width=setup_document(d,ref)
    intro,exp,tables,refs,order,mapping=prepared_markdown(src)
    eqs=[]
    for i,p in enumerate(original_paras):
        if i==0:format_para(p,'h1');continue
        if re.match(r'^3\.\d',p.text):
            for r in p.runs:
                if r.text:r.text=re.sub(r'^(3\.\d+)\s*',r'\1 ',r.text)
            format_para(p,'h2');continue
        format_para(p,'body')
        n=clean_numbered_equation(p,width)
        if n:eqs.append(n)
    assert eqs==list(range(1,33)),eqs
    interval=original_paras[45]._p.xpath('.//m:oMath')[0]
    cc=list(interval)
    for tt in xp(cc[1],'.//m:t'):tt.text=(tt.text or '').replace('∈(','∈')
    dd=OxmlElement('m:d');dp=OxmlElement('m:dPr')
    for tag,val in [('m:begChr','('),('m:endChr',']')]:
        ee=OxmlElement(tag);ee.set(qn('m:val'),val);dp.append(ee)
    dd.append(dp);inside=OxmlElement('m:e')
    for child in cc[2:-1]:inside.append(copy.deepcopy(child))
    dd.append(inside)
    for child in cc[2:]:interval.remove(child)
    interval.append(dd)
    fig=original_paras[7]
    caption=fig.text.replace('总体架构。','图 1 本文方法的总体架构。',1)
    pic=copy.deepcopy(fig._p)
    for child in list(pic):
        if child.tag!=qn('w:pPr') and not xp(child,'.//w:object'):pic.remove(child)
    picp=Paragraph(pic,d._body);format_para(picp,'figure_caption');picp.alignment=WD_ALIGN_PARAGRAPH.CENTER;picp.paragraph_format.keep_with_next=True
    ns={'v':'urn:schemas-microsoft-com:vml'}
    for shape in pic.xpath('.//v:shape',namespaces=ns) if type(pic)==etree._Element else etree.ElementTree(pic).xpath('.//v:shape',namespaces=ns):
        st=shape.get('style','');wm=re.search(r'width:([\d.]+)pt',st);hm=re.search(r'height:([\d.]+)pt',st)
        if wm and hm:
            ow,oh=float(wm[1]),float(hm[1]); nw=min(ow,width);nh=oh*nw/ow
            st=re.sub(r'width:[\d.]+pt',f'width:{nw:.3f}pt',st);st=re.sub(r'height:[\d.]+pt',f'height:{nh:.3f}pt',st);shape.set('style',st)
    fig._p.addprevious(pic);fig.clear();fig.add_run(caption);format_para(fig,'figure_caption')
    first=original_paras[5]
    for r in first.runs:
        if '总体结构。' in r.text:r.text=r.text.replace('总体结构。','总体结构。图 1 展示了本文方法的整体流程。',1);break
    for r in original_paras[82].runs:
        if '校准堆叠' in r.text:r.text=r.text.replace('校准堆叠',f'校准堆叠[{mapping[9]}]',1);break
    front=convert_md(intro,work,'introduction');back=convert_md(exp,work,'experiments')
    before=prepare_part(front,d,tables,width)
    after=prepare_part(back,d,tables,width,True)
    anchor=original_paras[0]._p
    frontmatter=[]
    for text,kind in [(TITLE,'title'),('摘要：'+ABSTRACT,'abstract'),('关键词：'+KEYWORDS,'keywords')]:
        p=d.add_paragraph(text);format_para(p,kind);d._body._body.remove(p._p);frontmatter.append(p._p)
        if kind in ['abstract','keywords']:
            p.clear();label,rest=text.split('：',1);r=p.add_run(label+'：');run_font(r);r.bold=True
            pieces=re.split(r'((?:VGGSound|UCF|ActivityNet)-GZSL)',rest)
            for piece in pieces:
                run_font(p.add_run(piece))
                if re.fullmatch(r'(?:VGGSound|UCF|ActivityNet)-GZSL',piece):
                    sr=p.add_run('cls');run_font(sr);sr.font.superscript=True
    for node in frontmatter+before:anchor.addprevious(node)
    sect=d._body._body.find(qn('w:sectPr'))
    for node in after:sect.addprevious(node)
    p=d.add_paragraph('5 结论');format_para(p,'h1')
    p=d.add_paragraph(CONCLUSION);format_para(p)
    p=d.add_paragraph('参考文献');format_para(p,'h1')
    for new,old in enumerate(order,1):
        p=d.add_paragraph(f'[{new}] '+refs[old]);format_para(p,'reference')
    for tag in ['w:commentRangeStart','w:commentRangeEnd','w:commentReference','w:proofErr']:
        for node in d._element.xpath('.//'+tag):node.getparent().remove(node)
    d.save(out);stable_zip(out)
    result=Document(out)
    numeric_source=[re.findall(r'\d+\.\d{2}',BeautifulSoup(h,'html.parser').get_text()) for h in tables]
    numeric_output=[re.findall(r'\d+\.\d{2}',' '.join(t._tbl.xpath('.//w:t/text()'))) for t in result.tables]
    assert numeric_source==numeric_output,'A numeric result changed during table conversion'
    assert len(result.tables)==9
    with ZipFile(src/'王虎-理论部分-0908.docx') as z:oldparts={n:z.read(n) for n in z.namelist() if n.startswith(('word/media/','word/embeddings/')) and not n.endswith('/')}
    with ZipFile(out) as z:
        for n,b in oldparts.items():assert z.read(n)==b,'Original figure/embedding changed'
        xml=z.read('word/document.xml')
    report={'output':out.name,'source_hashes':{f:hashlib.sha256((src/f).read_bytes()).hexdigest() for f in ['KA-GZAL中文.docx','王虎-理论部分-0908.docx','实验部分.md','引言与相关工作.md']},'data_tables':9,'theory_equations':eqs,'metric_equation':33,'references':len(order),'reference_mapping':mapping,'preserved_media_and_embeddings':list(oldparts),'result_values_unchanged':True,'file_sha256':hashlib.sha256(out.read_bytes()).hexdigest()}
    (work/'assembly_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--sources',required=True);ap.add_argument('--output',required=True);ap.add_argument('--work',required=True);a=ap.parse_args();build(a.sources,a.output,a.work)
